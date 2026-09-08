"""Company-isolated evidence store, reports, work, and optional notebook exchange."""

import base64
import hashlib
import io
import json
import re
import sqlite3
import time
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from businessos.finance import (ValidationError, decimal, historical_ltm, money,
                               normalize_period, project_cash, reconcile, required, scenario)

MAX_FILE_BYTES = 25 * 1024 * 1024
ID = re.compile(r"^[a-zA-Z0-9][a-zA-Z0-9_-]{0,95}$")
CHECKS = {
    "general": ["historical_financials", "tax_reconciliation", "bank_reconciliation", "debt_and_liens",
                "working_capital", "customer_concentration", "owner_replacement", "contracts_and_licenses"],
    "construction": ["job_costs_and_wip", "retainage_collectability", "backlog_cost_to_complete", "bonding_and_capacity"],
}


def now():
    return datetime.now(timezone.utc).isoformat()


def encode(value):
    return json.dumps(value, sort_keys=True, separators=(",", ":"), allow_nan=False)


def digest(value):
    return hashlib.sha256(value if isinstance(value, bytes) else encode(value).encode()).hexdigest()


def identifier(value):
    if not isinstance(value, str) or not ID.fullmatch(value):
        raise ValidationError("Invalid identifier")
    return value


def notebook_url(value):
    parsed = urlparse(value)
    match = re.fullmatch(r"/notebook/([a-f0-9-]{36})", parsed.path)
    if parsed.scheme != "https" or parsed.netloc not in ("notebook.google.com", "notebooklm.google.com") or not match or parsed.query or parsed.fragment:
        raise ValidationError("Use the exact HTTPS Google notebook URL without query parameters")
    try:
        uuid.UUID(match.group(1))
    except ValueError:
        raise ValidationError("Invalid notebook UUID") from None
    return value


def extract(data, filename):
    """Keep page/cell locators. Extraction is evidence, never verification."""
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".md", ".csv", ".json"):
        text = data.decode("utf-8-sig")
        return [{"locator": f"line:{n}", "text": line} for n, line in enumerate(text.splitlines(), 1) if line.strip()]
    if suffix == ".pdf":
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValidationError("Encrypted PDF needs an authorized decrypted copy")
        return [{"locator": f"page:{n}", "text": page.extract_text() or ""} for n, page in enumerate(reader.pages, 1)]
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        try:
            for sheet, cached in zip(formulas, values):
                for row, cached_row in zip(sheet.iter_rows(), cached.iter_rows()):
                    for cell, cached_cell in zip(row, cached_row):
                        if cell.value is not None:
                            rows.append({"locator": f"sheet:{sheet.title}!{cell.coordinate}", "text": str(cell.value),
                                         "cached_value": str(cached_cell.value) if cached_cell.value is not None else None,
                                         "formula": cell.data_type == "f"})
                            if len(rows) > 250000:
                                raise ValidationError("Workbook exceeds extraction cell limit; split the workbook")
        finally:
            formulas.close()
            values.close()
        return rows
    if suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(data))
        rows = [{"locator": f"paragraph:{n}", "text": p.text} for n, p in enumerate(document.paragraphs, 1) if p.text]
        for n, table in enumerate(document.tables, 1):
            for r, row in enumerate(table.rows, 1):
                for c, cell in enumerate(row.cells, 1):
                    rows.append({"locator": f"table:{n}/row:{r}/cell:{c}", "text": cell.text})
        return rows
    raise ValidationError("Supported source formats: PDF, XLSX, DOCX, TXT, MD, CSV, JSON")


class Runtime:
    def __init__(self, root, business_id):
        self.business_id = identifier(business_id)
        self.path = Path(root).resolve() / self.business_id
        self.path.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.blobs = self.path / "evidence"
        self.blobs.mkdir(exist_ok=True, mode=0o700)
        self.dbpath = self.path / "runtime.sqlite3"
        with self.db() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS records(kind TEXT, id TEXT, data TEXT NOT NULL, PRIMARY KEY(kind,id));
                CREATE TABLE IF NOT EXISTS events(seq INTEGER PRIMARY KEY AUTOINCREMENT, at TEXT, actor TEXT, action TEXT, record_id TEXT);
                CREATE TABLE IF NOT EXISTS cache(key TEXT PRIMARY KEY, data TEXT NOT NULL);
            """)
        self.dbpath.chmod(0o600)

    @contextmanager
    def db(self):
        db = sqlite3.connect(self.dbpath, timeout=15)
        try:
            db.execute("PRAGMA foreign_keys=ON")
            with db:
                yield db
        finally:
            db.close()

    def _get(self, db, kind, record_id):
        row = db.execute("SELECT data FROM records WHERE kind=? AND id=?", (kind, record_id)).fetchone()
        if row is None:
            raise ValidationError(f"{kind} record not found")
        return json.loads(row[0])

    def get(self, kind, record_id):
        with self.db() as db:
            return self._get(db, kind, record_id)

    def list(self, kind):
        with self.db() as db:
            return [json.loads(r[0]) for r in db.execute("SELECT data FROM records WHERE kind=? ORDER BY rowid DESC", (kind,))]

    def _put(self, db, kind, record_id, data, actor, action):
        db.execute("INSERT INTO records VALUES(?,?,?) ON CONFLICT(kind,id) DO UPDATE SET data=excluded.data", (kind, record_id, encode(data)))
        db.execute("INSERT INTO events(at,actor,action,record_id) VALUES(?,?,?,?)", (now(), actor, action, record_id))

    def configure(self, data, actor):
        required(data, ("display_name", "industry"))
        if data.get("notebook_url"):
            notebook_url(data["notebook_url"])
        clean = {"display_name": str(data["display_name"])[:160], "industry": data["industry"],
                 "notebook_url": data.get("notebook_url"), "updated_at": now(), "business_id": self.business_id}
        with self.db() as db:
            self._put(db, "config", "company", clean, actor, "company.configured")
        return clean

    def ingest(self, data, filename, document_id, origin, actor):
        identifier(document_id)
        if origin not in ("primary", "seller_statement", "external", "notebook", "businessos"):
            raise ValidationError("Invalid source origin")
        if not data or len(data) > MAX_FILE_BYTES:
            raise ValidationError("Source must contain 1 byte to 25 MB")
        # Prevent compressed document bombs before calling an XML parser.
        if Path(filename).suffix.lower() in (".xlsx", ".docx"):
            import zipfile
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValidationError("Expanded document exceeds 100 MB")
        sha = digest(data)
        source_id = f"{document_id}-{sha[:16]}"
        filename = Path(filename.replace("\\", "/")).name
        cache_key = f"extract-v1:{sha}:{Path(filename).suffix.lower()}"
        with self.db() as db:
            existing = db.execute("SELECT data FROM records WHERE kind='source' AND id=?", (source_id,)).fetchone()
            if existing:
                return {**json.loads(existing[0]), "duplicate": True}
            cached = db.execute("SELECT data FROM cache WHERE key=?", (cache_key,)).fetchone()
        started = time.monotonic()
        segments = json.loads(cached[0]) if cached else extract(data, filename)
        destination = self.blobs / sha
        if not destination.exists():
            try:
                with destination.open("xb") as stream:
                    stream.write(data)
            except FileExistsError:
                pass
            destination.chmod(0o600)
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            duplicate = db.execute("SELECT data FROM records WHERE kind='source' AND id=?", (source_id,)).fetchone()
            if duplicate:
                return {**json.loads(duplicate[0]), "duplicate": True}
            previous = [json.loads(r[0]) for r in db.execute("SELECT data FROM records WHERE kind='source'")]
            if any(r["sha256"] == sha and r["origin"] != origin for r in previous):
                raise ValidationError("Identical content cannot be relabeled as a different source origin")
            versions = [r for r in previous if r["document_id"] == document_id]
            for older in versions:
                if older["origin"] != origin:
                    raise ValidationError("A document's origin cannot change across versions")
                older["superseded"] = True
                self._put(db, "source", older["id"], older, actor, "source.superseded")
            record = {"id": source_id, "document_id": document_id, "version": len(versions) + 1, "sha256": sha,
                      "filename": filename, "origin": origin, "retrieved_at": now(), "superseded": False,
                      "segments": segments, "extraction_status": "extracted" if any(s["text"].strip() for s in segments) else "needs_ocr",
                      "independent_evidence": origin in ("primary", "external"),
                      "duration_ms": round((time.monotonic() - started) * 1000), "cache_hit": bool(cached),
                      "verification": "unreviewed"}
            db.execute("INSERT OR IGNORE INTO cache VALUES(?,?)", (cache_key, encode(segments)))
            self._put(db, "source", source_id, record, actor, "source.ingested")
        return record

    def source_bytes(self, source_id):
        source = self.get("source", source_id)
        path = self.blobs / source["sha256"]
        data = path.read_bytes()
        if digest(data) != source["sha256"]:
            raise ValidationError("Evidence integrity check failed")
        return source, data

    def source_refs(self, ids, independent=False):
        if not isinstance(ids, list) or not ids or len(ids) != len(set(ids)):
            raise ValidationError("Nonempty unique source IDs required")
        for source_id in ids:
            source = self.get("source", source_id)
            if source["superseded"]:
                raise ValidationError("Source version is superseded; review the updated source")
            if source["extraction_status"] != "extracted":
                raise ValidationError("Source needs readable text or OCR")
            if independent and not source["independent_evidence"]:
                raise ValidationError("Notebook, seller claims, and BusinessOS reports are not independent corroboration")

    def create_report(self, data, actor):
        from jsonschema import Draft202012Validator, FormatChecker
        schema_path = Path(__file__).resolve().parents[1] / "schemas" / "underwriting-input.v2.schema.json"
        errors = list(Draft202012Validator(json.loads(schema_path.read_text()), format_checker=FormatChecker()).iter_errors(data))
        if errors:
            raise ValidationError(f"{errors[0].json_path}: {errors[0].message}")
        required(data, ("report_id", "as_of_date", "periods", "adjustments", "reconciliations", "scenarios", "coverage", "citations", "recommendation"))
        identifier(data["report_id"])
        from businessos.finance import date_value
        date_value(data["as_of_date"])
        if data.get("business_id", self.business_id) != self.business_id:
            raise ValidationError("Cross-company report rejected")
        if data.get("schema_version", 2) != 2:
            raise ValidationError("Legacy report must be explicitly migrated before promotion")
        recommendation = data["recommendation"]
        required(recommendation, ("decision", "summary", "conditions"))
        if recommendation["decision"] not in ("pause", "decline", "proceed", "proceed_with_conditions"):
            raise ValidationError("Invalid recommendation")
        for collection in ("periods", "adjustments", "reconciliations", "scenarios", "coverage", "citations"):
            if not isinstance(data[collection], list):
                raise ValidationError(f"{collection} must be an array")
        period_ids = [p.get("id") for p in data["periods"]]
        if len(period_ids) != len(set(period_ids)):
            raise ValidationError("Duplicate period IDs")
        for adjustment in data["adjustments"]:
            if adjustment.get("period_id") not in period_ids:
                raise ValidationError("Adjustment refers to unknown period")
        periods = [normalize_period(p, data["adjustments"]) for p in data["periods"]]
        if any(p["status"] == "actual" and p["end"] > data["as_of_date"] for p in periods):
            raise ValidationError("Actual periods cannot end after the report date")
        if data.get("ltm_period_ids") and (len(set(data["ltm_period_ids"])) != len(data["ltm_period_ids"]) or set(data["ltm_period_ids"]) - set(period_ids)):
            raise ValidationError("LTM references must be unique known period IDs")
        for row in periods + data["adjustments"]:
            self.source_refs(row["source_ids"])
        reconciliations = [reconcile(r) for r in data["reconciliations"]]
        for row in reconciliations:
            self.source_refs([row["source_a"], row["source_b"]])
            self.source_refs([row["source_b"]], independent=True)
        citations = data["citations"]
        for citation in citations:
            required(citation, ("source_id", "locator", "claim"))
            source = self.get("source", citation["source_id"])
            self.source_refs([source["id"]])
            if citation["locator"] not in {s["locator"] for s in source["segments"]}:
                raise ValidationError("Citation locator does not exist in the referenced version")
        cited = {c["source_id"] for c in citations}
        used = {s for p in periods + data["adjustments"] for s in p["source_ids"]}
        used |= {r[k] for r in reconciliations for k in ("source_a", "source_b")}
        if used - cited:
            raise ValidationError("All financial inputs require page/cell/line citations")
        config = self.get("config", "company")
        checks = CHECKS["general"] + (CHECKS["construction"] if config["industry"] == "construction" else [])
        coverage = {}
        for check in data["coverage"]:
            required(check, ("id", "status", "notes"))
            if check["id"] in coverage or check["id"] not in checks:
                raise ValidationError("Duplicate or unknown coverage item")
            if check["status"] not in ("missing", "in_progress", "complete"):
                raise ValidationError("Coverage status must be missing, in_progress, or complete")
            if check["status"] == "complete":
                self.source_refs(check.get("source_ids"), independent=True)
                if set(check["source_ids"]) - cited:
                    raise ValidationError("Completed coverage requires precise citations")
            coverage[check["id"]] = check
        gaps = [name for name in checks if coverage.get(name, {}).get("status") != "complete"]
        coverage = [coverage.get(name, {"id": name, "status": "missing", "notes": "Evidence not established"}) for name in checks]
        names = [s["name"] for s in data["scenarios"]]
        if len(names) != len(set(names)):
            raise ValidationError("Duplicate scenario names")
        results = []
        for row in data["scenarios"]:
            if row["name"] not in ("base", "downside", "severe_downside", "upside"):
                raise ValidationError("Unknown scenario name")
            required(row, ("assumptions", "source_ids"))
            self.source_refs(row["source_ids"])
            if set(row["source_ids"]) - cited:
                raise ValidationError("Scenario source inputs need precise citations")
            if not row["assumptions"]:
                raise ValidationError("Scenario assumptions must be disclosed")
            results.append(scenario(row))
        if not {"base", "downside", "severe_downside"}.issubset(names):
            gaps.append("three_downside_scenarios")
        if not periods:
            gaps.append("financial_periods")
        if not reconciliations:
            gaps.append("financial_reconciliations")
        gaps.extend(f"reconciliation:{r['id']}" for r in reconciliations if r["status"] in ("open", "waived"))
        ltm = historical_ltm([p for p in periods if p["id"] in data["ltm_period_ids"]]) if data.get("ltm_period_ids") else None
        if not ltm:
            gaps.append("actual_ltm")
        actual_fiscal = [p for p in periods if p["status"] == "actual" and p.get("period_type") == "fiscal_year"]
        fiscal_ranges = sorted((date_value(p["start"]), date_value(p["end"])) for p in actual_fiscal)
        for index, (start, end) in enumerate(fiscal_ranges):
            if (end - start).days not in (364, 365) or (index and start <= fiscal_ranges[index - 1][1]):
                raise ValidationError("Fiscal years must be complete and non-overlapping")
        if len(actual_fiscal) < 3:
            gaps.append("three_fiscal_years")
        if gaps and recommendation["decision"] == "proceed":
            raise ValidationError("Unconditional proceed is not permitted with material gaps")
        if gaps and recommendation["decision"] == "proceed_with_conditions" and not recommendation["conditions"]:
            raise ValidationError("Conditional recommendation must disclose conditions")
        fingerprint = digest(data)
        record = {"id": data["report_id"], "schema_version": 2, "business_id": self.business_id,
                  "created_at": now(), "created_by": actor, "as_of_date": data["as_of_date"],
                  "input_hash": fingerprint, "status": "working", "periods": periods, "ltm": ltm,
                  "reconciliations": reconciliations, "scenarios": results, "coverage": coverage,
                  "gaps": gaps, "citations": citations, "recommendation": recommendation, "inputs": data,
                  "source_ids": sorted(cited | used), "risks": data.get("risks", []), "review": None}
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            existing = db.execute("SELECT data FROM records WHERE kind='report' AND id=?", (record["id"],)).fetchone()
            if existing:
                existing = json.loads(existing[0])
                if existing["input_hash"] == fingerprint:
                    return {**existing, "cache_hit": True}
                raise ValidationError("Reports are immutable; use a new report ID for revised inputs")
            self._put(db, "report", record["id"], record, actor, "report.calculated")
        return record

    def review_report(self, report_id, decision, exceptions, rationale, actor):
        if decision not in ("approve", "reject") or not rationale.strip():
            raise ValidationError("Review requires approve/reject and rationale")
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            report = self._get(db, "report", report_id)
            if actor == report["created_by"]:
                raise ValidationError("A report author cannot approve their own work")
            if decision == "approve":
                for source_id in report["source_ids"]:
                    source = self._get(db, "source", source_id)
                    if source["superseded"]:
                        raise ValidationError("Evidence changed; generate and review a new report")
                if set(exceptions) != set(report["gaps"]):
                    raise ValidationError("Review must explicitly acknowledge every material gap")
            report["review"] = {"reviewer": actor, "at": now(), "decision": decision, "exceptions": exceptions, "rationale": rationale, "input_hash": report["input_hash"]}
            report["status"] = ("decision_ready_with_exceptions" if exceptions else "final") if decision == "approve" else "working"
            self._put(db, "report", report_id, report, actor, "report.reviewed")
        return report

    def memo(self, report_id):
        r = self.get("report", report_id)
        lines = [f"# {self.get('config', 'company')['display_name']} — underwriting", "", f"As of {r['as_of_date']} | {r['status']} | Report {r['id']}", "",
                 f"## Recommendation: {r['recommendation']['decision']}", r["recommendation"]["summary"], "",
                 "## Conditions and unresolved evidence"]
        lines += [f"- {item}" for item in r["recommendation"]["conditions"] + r["gaps"]] or ["No recorded gaps; independent review remains required."]
        lines += ["", "## Financial periods", "| Period | Classification | Definition | Revenue | Normalized earnings |", "|---|---|---|---:|---:|"]
        lines += [f"| {p['start']} to {p['end']} | {p['status']} | {p['earnings_definition']} | {p['revenue']} | {p['normalized_earnings']} |" for p in r["periods"]]
        lines += ["", "## Scenario assumptions and results"]
        for s in r["scenarios"]:
            lines += [f"### {s['name']}", f"EBITDA: {s['normalized_ebitda']}; ending cash: {s['ending_cash']}; funding gap: {s['funding_gap']}; enterprise value: {s['enterprise_value'] or 'not established'}."]
            lines += [f"- {a}" for a in s["inputs"]["assumptions"]]
        lines += ["", "## Risks and mitigations"]
        lines += [f"- {risk.get('title', 'Risk')}: {risk.get('mitigation', 'Not established')}" for risk in r["risks"]]
        lines += ["", "## Source-linked claims"]
        lines += [f"- {c['claim']} [{c['source_id']} / {c['locator']}]" for c in r["citations"]]
        lines += ["", "BusinessOS-generated analysis. Do not count this report as independent corroboration of its own inputs."]
        return "\n".join(lines)

    def save_projection(self, data, actor):
        required(data, ("opening_cash", "months", "initiatives", "capacity"))
        results = project_cash(data["opening_cash"], data["months"], data["initiatives"], data["capacity"])
        key = digest(data)
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            if not db.execute("SELECT 1 FROM records WHERE kind='projection' AND id=?", (key,)).fetchone():
                self._put(db, "projection", key, {"id": key, "inputs": data, "results": results, "created_at": now(), "created_by": actor}, actor, "projection.calculated")
        return results

    def save_initiative(self, data, actor):
        required(data, ("id", "title", "owner", "baseline", "target", "metric", "measurement_end", "stop_condition", "rollback", "source_ids", "benefit_group", "setup_cost", "monthly_cost", "monthly_cash_benefit", "monthly_hours", "start_month", "ramp_months"))
        identifier(data["id"])
        self.source_refs(data["source_ids"])
        from businessos.finance import date_value
        date_value(data["measurement_end"])
        decimal(data["baseline"])
        decimal(data["target"])
        project_cash(0, [], [data], {"monthly_hours": 0, "minimum_cash": 0})
        record = {**data, "state": "proposed", "created_by": actor, "created_at": now(), "observations": []}
        with self.db() as db:
            if db.execute("SELECT 1 FROM records WHERE kind='initiative' AND id=?", (data["id"],)).fetchone():
                raise ValidationError("Initiative exists; create a revised initiative with a new ID")
            self._put(db, "initiative", data["id"], record, actor, "initiative.proposed")
        return record

    def initiative_action(self, initiative_id, data, actor, stage):
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            item = self._get(db, "initiative", initiative_id)
            action = data.get("action")
            if action == "observe":
                required(data, ("value", "source_id", "as_of_date"))
                self.source_refs([data["source_id"]])
                from businessos.finance import date_value
                date_value(data["as_of_date"])
                item["observations"].append({"value": money(data["value"]), "source_id": data["source_id"], "as_of_date": data["as_of_date"], "recorded_by": actor})
            else:
                states = {"proposed": ("shadow", "stopped"), "shadow": ("pilot", "stopped"), "pilot": ("active", "stopped"), "active": ("closed", "stopped"), "stopped": (), "closed": ()}
                if action not in states[item["state"]]:
                    raise ValidationError("Invalid initiative transition")
                required(data, ("rationale",))
                if action in ("pilot", "active"):
                    if stage not in ("transition", "operating") or actor == item["created_by"]:
                        raise ValidationError("Pilot/live promotion requires operating authority and independent approval")
                    if not item["observations"]:
                        raise ValidationError("Pilot/live promotion requires measured evidence")
                item["state"] = action
                item["decision"] = {"actor": actor, "at": now(), "rationale": data["rationale"]}
            self._put(db, "initiative", initiative_id, item, actor, "initiative." + str(action))
        return item

    def notebook_job(self, data, actor):
        config = self.get("config", "company")
        url = notebook_url(config.get("notebook_url", ""))
        action = data.get("action")
        if action not in ("read", "ask", "publish"):
            raise ValidationError("Notebook action must be read, ask, or publish")
        payload = {"action": action, "notebook_url": url}
        if action == "ask":
            required(data, ("question",))
            payload["question"] = str(data["question"])[:10000]
        if action == "publish":
            report = self.get("report", data.get("report_id", ""))
            if not report["review"] or report["review"]["decision"] != "approve":
                raise ValidationError("Only independently reviewed reports may be published")
            self.source_refs(report["source_ids"])
            payload.update({"report_id": report["id"], "content": self.memo(report["id"]), "report_hash": report["input_hash"]})
        key = digest(payload)
        job_id = key[:32] if action == "publish" else uuid.uuid4().hex
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            old = db.execute("SELECT data FROM records WHERE kind='notebook_job' AND id=?", (job_id,)).fetchone()
            if old:
                return {**json.loads(old[0]), "duplicate": True}
            record = {"id": job_id, **payload, "status": "pending", "requested_at": now(), "requested_by": actor,
                      "content_hash": key, "result": None, "attempts": 0}
            self._put(db, "notebook_job", job_id, record, actor, "notebook.requested")
        return record

    def notebook_result(self, job_id, data, actor):
        """Browser worker reports explicit evidence; uncertain uploads never retry automatically."""
        if data.get("status") not in ("verified", "login_required", "failed", "uncertain"):
            raise ValidationError("Invalid notebook result status")
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            job = self._get(db, "notebook_job", job_id)
            if job["status"] == "verified":
                return job
            if data.get("notebook_url") != job["notebook_url"]:
                raise ValidationError("Notebook identity mismatch")
            if data["status"] == "verified":
                required(data, ("observed_at", "evidence", "sources"))
                if not isinstance(data["sources"], list):
                    raise ValidationError("Notebook source inventory must be an array")
                if job["action"] == "publish":
                    required(data, ("source_id", "content_hash", "processing_complete"))
                    if data["content_hash"] != job["content_hash"] or data["processing_complete"] is not True:
                        raise ValidationError("Publication requires matching content and completed processing")
                if job["action"] == "ask":
                    required(data, ("answer", "citations"))
                    if not isinstance(data["citations"], list) or not data["citations"]:
                        raise ValidationError("Notebook answer needs captured citations")
            job.update({"status": data["status"], "result": data, "finished_at": now(), "attempts": job["attempts"] + 1})
            self._put(db, "notebook_job", job_id, job, actor, "notebook." + data["status"])
        return job

    def snapshot(self):
        # Derive health from persisted observations; do not call Notebook here.
        try:
            config = self.get("config", "company")
        except ValidationError:
            config = {"business_id": self.business_id, "display_name": self.business_id, "industry": "unknown", "notebook_url": None}
        sources = self.list("source")
        reports = self.list("report")
        current_ids = {s["id"] for s in sources if not s["superseded"]}
        for report in reports:
            report["stale"] = bool(set(report["source_ids"]) - current_ids)
        jobs = self.list("notebook_job")
        with self.db() as db:
            events = [dict(zip(("at", "actor", "action", "record_id"), r)) for r in db.execute("SELECT at,actor,action,record_id FROM events ORDER BY seq DESC LIMIT 30")]
        return {"business": config, "sources": [{k: v for k, v in s.items() if k != "segments"} for s in sources],
                "reports": reports, "initiatives": self.list("initiative"), "projections": self.list("projection"), "notebook_jobs": jobs,
                "events": events, "generated_at": now(), "mode": "live",
                "integration_health": {"evidence_store": "healthy", "notebook": jobs[0]["status"] if jobs else "disconnected",
                                       "accounting": "not_connected", "banking": "not_connected"},
                "processing": {"sources": len(sources), "cache_hits": sum(int(s["cache_hit"]) for s in sources),
                               "extraction_ms": sum(s["duration_ms"] for s in sources), "model_cost_usd": "0.00"}}

    def claim_notebook_job(self, job_id, actor):
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            job = self._get(db, "notebook_job", job_id)
            if job["status"] != "pending":
                raise ValidationError("Job is already claimed or finished; inspect it before any retry")
            job.update({"status": "running", "claimed_at": now(), "claimed_by": actor})
            self._put(db, "notebook_job", job_id, job, actor, "notebook.claimed")
        return job
